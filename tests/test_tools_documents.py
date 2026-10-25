import asyncio

from openpyxl import Workbook

from tests.fakes import make_registry


def make_pdf(path, pages: list[str]) -> None:
    n = len(pages)
    objs: list[bytes] = []
    kids = " ".join(f"{3 + i} 0 R" for i in range(n))
    objs.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    objs.append(f"<< /Type /Pages /Kids [{kids}] /Count {n} >>".encode())
    font_id = 3 + 2 * n
    for i, text in enumerate(pages):
        content_id = 3 + n + i
        objs.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 {font_id} 0 R >> >> /Contents {content_id} 0 R >>".encode()
        )
    for text in pages:
        stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode("latin-1")
        objs.append(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")
    objs.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objs, 1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n".encode() + b"0000000000 65535 f \n"
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    path.write_bytes(bytes(out))


def runner(tmp_path):
    reg = make_registry(tmp_path)
    return reg, lambda name, **a: asyncio.run(reg.run(name, a))


def test_read_pdf_pages(tmp_path):
    _, run = runner(tmp_path)
    make_pdf(tmp_path / "r.pdf", ["Birinci sayfa", "Second page", "Third"])
    out = run("read_pdf", filename="r.pdf")
    assert out.startswith("Pages 1-3 of 3:") and "Second page" in out
    assert run("read_pdf", filename="r.pdf", start_page=2, end_page=2) == "Pages 2-2 of 3:\n\nSecond page"
    assert run("read_pdf", filename="r.pdf", start_page=5) == "Invalid page range."
    assert run("read_pdf", filename="r.pdf", start_page=3, end_page=2) == "Invalid page range."
    assert run("read_pdf", filename="r.pdf", start_page=2, end_page=99).startswith("Pages 2-3 of 3:")


def test_read_pdf_empty_page_and_errors(tmp_path):
    _, run = runner(tmp_path)
    make_pdf(tmp_path / "e.pdf", ["Hello", ""])
    assert "[page 2: no extractable text]" in run("read_pdf", filename="e.pdf")
    assert run("read_pdf", filename="nope.pdf") == "File 'nope.pdf' not found."
    assert run("read_pdf", filename="../x.pdf") == "Path is outside the workspace."
    (tmp_path / "bad.pdf").write_bytes(b"not a pdf")
    assert run("read_pdf", filename="bad.pdf") == "Failed to read the PDF."


def test_read_pdf_page_cap(tmp_path):
    _, run = runner(tmp_path)
    make_pdf(tmp_path / "big.pdf", [f"p{i}" for i in range(40)])
    assert run("read_pdf", filename="big.pdf").startswith("Pages 1-30 of 40:")


def test_spreadsheet_roundtrip_xlsx(tmp_path):
    _, run = runner(tmp_path)
    assert run("write_spreadsheet", filename="t.xlsx", rows=[["Ad", "Tutar"], ["Çay", 12.5], ["Şeker", None]]) == "Wrote 3 rows to 't.xlsx'."
    out = run("read_spreadsheet", filename="t.xlsx")
    assert "Sheets: Sayfa1" in out and "Sayfa1: 3 rows x 2 columns" in out
    assert "| Ad | Tutar |" in out and "| Çay | 12.5 |" in out and "| Şeker |  |" in out
    assert "|---|---|" in out


def test_spreadsheet_pipe_escape_and_csv_write(tmp_path):
    _, run = runner(tmp_path)
    assert run("write_spreadsheet", filename="p.csv", rows=[["a|b", "c"], [1, None]]) == "Wrote 2 rows to 'p.csv'."
    assert (tmp_path / "p.csv").read_bytes().startswith(b"\xef\xbb\xbf")
    assert r"| a\|b | c |" in run("read_spreadsheet", filename="p.csv")


def test_read_csv_semicolon_and_bom(tmp_path):
    _, run = runner(tmp_path)
    (tmp_path / "d.csv").write_text("\ufeffad;yaş\nAli;30\n", encoding="utf-8")
    assert "| ad | yaş |" in run("read_spreadsheet", filename="d.csv")


def test_read_spreadsheet_max_rows_and_sheet(tmp_path):
    _, run = runner(tmp_path)
    wb = Workbook()
    wb.active.title = "A"
    wb.active.append(["x"])
    b = wb.create_sheet("B")
    for i in range(300):
        b.append([f"r{i}", i])
    wb.save(tmp_path / "m.xlsx")
    wb.close()
    out = run("read_spreadsheet", filename="m.xlsx", sheet="B", max_rows=500)
    assert "B: 300 rows" in out and out.count("\n| ") <= 201
    assert run("read_spreadsheet", filename="m.xlsx", sheet="Z") == "Sheet 'Z' not found. Sheets: A, B"


def test_read_spreadsheet_errors(tmp_path):
    _, run = runner(tmp_path)
    assert run("read_spreadsheet", filename="nope.xlsx") == "File 'nope.xlsx' not found."
    (tmp_path / "x.txt").write_text("a")
    assert run("read_spreadsheet", filename="x.txt") == "Use a .xlsx or .csv file."
    assert run("read_spreadsheet", filename="../x.csv") == "Path is outside the workspace."


def test_write_spreadsheet_bad_extension(tmp_path):
    _, run = runner(tmp_path)
    assert run("write_spreadsheet", filename="t.ods", rows=[["a"]]) == "Use a .xlsx or .csv file name."


def test_write_spreadsheet_confirmation(tmp_path):
    reg, run = runner(tmp_path)
    assert reg.needs_confirmation("write_spreadsheet", {"filename": "t.csv", "rows": [["b"]]}) is None
    run("write_spreadsheet", filename="t.csv", rows=[["a"]])
    c = reg.needs_confirmation("write_spreadsheet", {"filename": "t.csv", "rows": [["b", 2], ["c", None]]})
    assert c.description == "‘t.csv’ üzerine yazılacak."
    assert c.preview == "b | 2\nc | "
    big = reg.needs_confirmation("write_spreadsheet", {"filename": "t.csv", "rows": [["x" * 500]] * 5})
    assert len(big.preview) == 600


def test_read_large_csv_streams_rows(tmp_path):
    _, run = runner(tmp_path)
    lines = ["a,b,c"] + [f"r{i},{i}" for i in range(1000)]
    (tmp_path / "big.csv").write_text("\n".join(lines), encoding="utf-8")
    out = run("read_spreadsheet", filename="big.csv", max_rows=500)
    assert "big: 1001 rows x 3 columns" in out
    table = [l for l in out.splitlines() if l.startswith("| ")]
    assert len(table) == 200
    assert "| r198 | 198 |" in out and "r300" not in out


def test_csv_cp1254_fallback(tmp_path):
    (tmp_path / "t.csv").write_bytes("ad;şehir\nİlhan;Ağrı\n".encode("cp1254"))
    out = asyncio.run(make_registry(tmp_path).run("read_spreadsheet", {"filename": "t.csv"}))
    assert "şehir" in out and "İlhan" in out and "Ağrı" in out


def test_write_spreadsheet_formula_cells_are_text(tmp_path):
    reg = make_registry(tmp_path)
    rows = [["=1+1", "+x", "-y", "@z", "ok", 5, -3]]
    asyncio.run(reg.run("write_spreadsheet", {"filename": "f.csv", "rows": rows}))
    text = (tmp_path / "f.csv").read_text(encoding="utf-8-sig")
    assert text.strip() == "'=1+1,'+x,'-y,'@z,ok,5,-3"
    asyncio.run(reg.run("write_spreadsheet", {"filename": "f.xlsx", "rows": rows}))
    from openpyxl import load_workbook

    ws = load_workbook(tmp_path / "f.xlsx").active
    cells = list(ws[1])
    assert [c.data_type for c in cells[:4]] == ["s"] * 4
    assert cells[0].value == "=1+1" and cells[5].value == 5 and cells[6].value == -3


def test_document_arg_coercion(tmp_path):
    (tmp_path / "t.csv").write_text("a\n1\n2\n")
    reg = make_registry(tmp_path)
    out = asyncio.run(reg.run("read_spreadsheet", {"filename": "t.csv", "sheet": None, "max_rows": None}))
    assert "3 rows" in out
    out = asyncio.run(reg.run("read_spreadsheet", {"filename": "t.csv", "max_rows": "1"}))
    assert "| a |" in out and "| 1 |" not in out
