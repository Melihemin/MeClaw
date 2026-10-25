"""Document tools: PDF reading and spreadsheet (.xlsx/.csv) reading and writing."""
from __future__ import annotations

import csv
import io

from meclaw.tools.base import Confirmation, ToolContext, ToolSpec

MAX_PDF_PAGES = 30
DEFAULT_SHEET_ROWS = 50
MAX_SHEET_ROWS = 200
SHEET_PREVIEW_ROWS = 20
WRITE_PREVIEW_CHARS = 600


def _obj(properties: dict, required: list[str] | None = None) -> dict:
    schema: dict = {"type": "object", "properties": properties}
    if required:
        schema["required"] = required
    return schema


_S = {"type": "string"}


def read_pdf(ctx: ToolContext, filename: str, start_page: int | None = None, end_page: int | None = None) -> str:
    from pypdf import PdfReader

    if not filename:
        return "No filename provided."
    try:
        path = ctx.explorer.resolve(filename)
    except ValueError:
        return "Path is outside the workspace."
    if not path.is_file():
        return f"File '{filename}' not found."
    try:
        reader = PdfReader(str(path))
        total = len(reader.pages)
        start = int(start_page) if start_page is not None else 1
        end = int(end_page) if end_page is not None else total
        if start < 1 or start > total or end < start:
            return "Invalid page range."
        end = min(end, total, start + MAX_PDF_PAGES - 1)
        parts = []
        for k in range(start, end + 1):
            text = (reader.pages[k - 1].extract_text() or "").strip()
            parts.append(text if text else f"[page {k}: no extractable text]")
    except Exception:
        return "Failed to read the PDF."
    return f"Pages {start}-{end} of {total}:\n\n" + "\n\n".join(parts)


def _cell(value) -> str:
    return "" if value is None else str(value).replace("|", "\\|").replace("\n", " ")


def _table(rows: list[list]) -> str:
    width = max(len(r) for r in rows)
    lines = []
    for i, row in enumerate(rows):
        cells = [_cell(c) for c in row] + [""] * (width - len(row))
        lines.append("| " + " | ".join(cells) + " |")
        if i == 0:
            lines.append("|" + "---|" * width)
    return "\n".join(lines)


def _read_csv(path) -> tuple[str, list[list[str]], int, int]:
    text = path.read_text(encoding="utf-8-sig")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = list(csv.reader(io.StringIO(text), dialect))
    return path.stem, rows, len(rows), max((len(r) for r in rows), default=0)


def read_spreadsheet(ctx: ToolContext, filename: str, sheet: str = "", max_rows: int = DEFAULT_SHEET_ROWS) -> str:
    if not filename:
        return "No filename provided."
    try:
        path = ctx.explorer.resolve(filename)
    except ValueError:
        return "Path is outside the workspace."
    if not path.is_file():
        return f"File '{filename}' not found."
    limit = max(1, min(int(max_rows), MAX_SHEET_ROWS))
    suffix = path.suffix.lower()
    try:
        if suffix == ".csv":
            title, rows, count, columns = _read_csv(path)
            names = [title]
            head = rows[:limit]
        elif suffix == ".xlsx":
            from openpyxl import load_workbook

            wb = load_workbook(path, read_only=True, data_only=True)
            try:
                names = list(wb.sheetnames)
                if sheet and sheet not in names:
                    return f"Sheet '{sheet}' not found. Sheets: {', '.join(names)}"
                title = sheet or names[0]
                head, count, columns = [], 0, 0
                for row in wb[title].iter_rows(values_only=True):
                    count += 1
                    columns = max(columns, len(row))
                    if count <= limit:
                        head.append(list(row))
            finally:
                wb.close()
        else:
            return "Use a .xlsx or .csv file."
    except Exception:
        return "Failed to read the spreadsheet."
    summary = f"Sheets: {', '.join(names)}\n{title}: {count} rows x {columns} columns"
    if not head:
        return summary
    return f"{summary}\n\n{_table(head)}"


def write_spreadsheet(ctx: ToolContext, filename: str, rows: list, sheet: str = "Sayfa1") -> str:
    if not filename:
        return "No filename provided."
    suffix = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if suffix not in ("xlsx", "csv"):
        return "Use a .xlsx or .csv file name."
    try:
        path = ctx.explorer.resolve(filename)
    except ValueError:
        return "Path is outside the workspace."
    if not isinstance(rows, list) or not all(isinstance(r, list) for r in rows):
        return "rows must be a list of lists."
    try:
        path.parent.mkdir(parents=True, exist_ok=True)
        if suffix == "csv":
            with open(path, "w", encoding="utf-8-sig", newline="") as fh:
                writer = csv.writer(fh)
                for row in rows:
                    writer.writerow(["" if c is None else c for c in row])
        else:
            from openpyxl import Workbook

            wb = Workbook()
            try:
                ws = wb.active
                ws.title = (sheet or "Sayfa1")[:31]
                for row in rows:
                    ws.append(list(row))
                wb.save(path)
            finally:
                wb.close()
    except Exception:
        return "Failed to write the spreadsheet."
    return f"Wrote {len(rows)} rows to '{filename}'."


def _confirm_write_spreadsheet(ctx: ToolContext, a: dict):
    f = a.get("filename", "")
    if not ctx.explorer.resolve(f).is_file():
        return None
    rows = a.get("rows") or []
    lines = [
        " | ".join("" if c is None else str(c) for c in r)
        for r in rows[:SHEET_PREVIEW_ROWS]
        if isinstance(r, list)
    ]
    return Confirmation(f"‘{f}’ üzerine yazılacak.", "\n".join(lines)[:WRITE_PREVIEW_CHARS])


DOCUMENT_TOOLS: list[ToolSpec] = [
    ToolSpec(
        "read_pdf",
        "Read text from a PDF in the workspace (at most 30 pages per call; pages are 1-based, inclusive).",
        _obj(
            {
                "filename": _S,
                "start_page": {"type": "integer", "description": "First page (default 1)."},
                "end_page": {"type": "integer", "description": "Last page (default: last, capped at 30 pages)."},
            },
            ["filename"],
        ),
        read_pdf,
    ),
    ToolSpec(
        "read_spreadsheet",
        "Read a .xlsx or .csv file: sheet names, size, and the first rows as a Markdown table.",
        _obj(
            {
                "filename": _S,
                "sheet": {"type": "string", "description": "Sheet name (xlsx only; default: first sheet)."},
                "max_rows": {"type": "integer", "description": "Rows to show (default 50, max 200)."},
            },
            ["filename"],
        ),
        read_spreadsheet,
    ),
    ToolSpec(
        "write_spreadsheet",
        "Write rows to a .xlsx or .csv file (chosen by extension). Overwrites an existing file.",
        _obj(
            {
                "filename": _S,
                "rows": {
                    "type": "array",
                    "description": "Rows; each row is an array of string, number or null cells. The first row is the header.",
                    "items": {"type": "array", "items": {"type": ["string", "number", "null"]}},
                },
                "sheet": {"type": "string", "description": "Sheet name for xlsx (default Sayfa1)."},
            },
            ["filename", "rows"],
        ),
        write_spreadsheet,
        confirm=_confirm_write_spreadsheet,
    ),
]
