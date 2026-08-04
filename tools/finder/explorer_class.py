import logging
import os
import shutil
import zipfile
from pathlib import Path
from typing import Optional

from docx import Document
from docx.shared import Pt


class Explorer:
    """Simple class-based file and folder explorer for local operations."""

    def __init__(self, base_dir: Optional[str] = None) -> None:
        self.base_dir = Path(base_dir or os.path.join(os.path.dirname(__file__), "testdir")).resolve()
        self.base_dir.mkdir(parents=True, exist_ok=True)

    def _build_path(self, name: str) -> Path:
        return self.base_dir / name

    def create_file(self, filename: str) -> str:
        if not filename:
            return "No filename provided."

        try:
            path = self._build_path(filename)
            path.touch(exist_ok=True)
            return f"File '{filename}' created successfully."
        except Exception as exc:
            logging.error("File Creation Error: %s", exc)
            return "Failed to create the file."

    def delete_file(self, filename: str) -> str:
        if not filename:
            return "No filename provided."

        try:
            path = self._build_path(filename)
            path.unlink(missing_ok=True)
            return f"File '{filename}' deleted successfully."
        except Exception as exc:
            logging.error("File Deletion Error: %s", exc)
            return "Failed to delete the file."

    def copy_file(self, source_filename: str, destination_filename: str) -> str:
        if not source_filename or not destination_filename:
            return "Source or destination filename not provided."

        try:
            source_path = self._build_path(source_filename)
            destination_path = self._build_path(destination_filename)
            destination_path.write_text(source_path.read_text(encoding="utf-8"), encoding="utf-8")
            return f"File '{source_filename}' copied to '{destination_filename}' successfully."
        except FileNotFoundError:
            return f"File '{source_filename}' not found."
        except Exception as exc:
            logging.error("File Copy Error: %s", exc)
            return "Failed to copy the file."

    def read_file(self, filename: str) -> str:
        if not filename:
            return "No filename provided."

        try:
            return self._build_path(filename).read_text(encoding="utf-8")
        except FileNotFoundError:
            return f"File '{filename}' not found."
        except Exception as exc:
            logging.error("File Read Error: %s", exc)
            return "Failed to read the file."

    def write_file(self, filename: str, content: str) -> str:
        if not filename:
            return "No filename provided."

        try:
            self._build_path(filename).write_text(content, encoding="utf-8")
            return f"Content written to '{filename}' successfully."
        except Exception as exc:
            logging.error("File Write Error: %s", exc)
            return "Failed to write to the file."

    def rename_file(self, old_filename: str, new_filename: str) -> str:
        if not old_filename or not new_filename:
            return "Old or new filename not provided."

        try:
            old_path = self._build_path(old_filename)
            new_path = self._build_path(new_filename)
            old_path.rename(new_path)
            return f"File '{old_filename}' renamed to '{new_filename}' successfully."
        except FileNotFoundError:
            return f"File '{old_filename}' not found."
        except Exception as exc:
            logging.error("File Rename Error: %s", exc)
            return "Failed to rename the file."

    def read_word_file(self, filename: str) -> str:
        if not filename:
            return "No filename provided."

        try:
            doc = Document(self._build_path(filename))
            return "\n".join(para.text for para in doc.paragraphs)
        except FileNotFoundError:
            return f"File '{filename}' not found."
        except Exception as exc:
            logging.error("Word File Read Error: %s", exc)
            return "Failed to read the Word file."

    def write_word_file(
        self,
        filename: str,
        content: str,
        font_name: Optional[str] = None,
        font_size: Optional[int] = None,
    ) -> str:
        if not filename:
            return "No filename provided."

        try:
            doc = Document()
            for line in content.splitlines():
                paragraph = doc.add_paragraph()
                run = paragraph.add_run(line)
                if font_name:
                    run.font.name = font_name
                if font_size:
                    run.font.size = Pt(font_size)
            doc.save(self._build_path(filename))
            return f"Content written to '{filename}' successfully."
        except Exception as exc:
            logging.error("Word File Write Error: %s", exc)
            return "Failed to write to the Word file."

    def create_folder(self, folder_name: str) -> str:
        if not folder_name:
            return "No folder name provided."

        try:
            folder_path = self._build_path(folder_name)
            folder_path.mkdir(parents=True, exist_ok=True)
            return f"Folder '{folder_name}' created successfully."
        except Exception as exc:
            logging.error("Folder Creation Error: %s", exc)
            return "Failed to create the folder."

    def remove_folder(self, folder_name: str) -> str:
        if not folder_name:
            return "No folder name provided."

        try:
            folder_path = self._build_path(folder_name)
            folder_path.rmdir()
            return f"Folder '{folder_name}' removed successfully."
        except FileNotFoundError:
            return f"Folder '{folder_name}' not found."
        except OSError as exc:
            logging.error("Folder Removal Error: %s", exc)
            return "Failed to remove the folder. Ensure it's empty before deletion."

    def list_items(self) -> list[dict[str, str]]:
        try:
            items: list[dict[str, str]] = []
            for path in sorted(self.base_dir.iterdir()):
                item_type = "dir" if path.is_dir() else "file"
                items.append({"name": path.name, "type": item_type})
            return items
        except Exception as exc:
            logging.error("Indexing Error: %s", exc)
            return []

    def move_item(self, source_path: str, destination_path: str) -> str:
        if not source_path or not destination_path:
            return "Source or destination path not provided."

        try:
            source = self._build_path(source_path)
            destination = self._build_path(destination_path)
            destination.parent.mkdir(parents=True, exist_ok=True)
            source.rename(destination)
            return f"Item '{source_path}' moved to '{destination_path}' successfully."
        except FileNotFoundError:
            return f"Item '{source_path}' not found."
        except Exception as exc:
            logging.error("Move Error: %s", exc)
            return "Failed to move the item."

    def rename_folder(self, old_folder_name: str, new_folder_name: str) -> str:
        if not old_folder_name or not new_folder_name:
            return "Old or new folder name not provided."

        try:
            old_path = self._build_path(old_folder_name)
            new_path = self._build_path(new_folder_name)
            old_path.rename(new_path)
            return f"Folder '{old_folder_name}' renamed to '{new_folder_name}' successfully."
        except FileNotFoundError:
            return f"Folder '{old_folder_name}' not found."
        except Exception as exc:
            logging.error("Folder Rename Error: %s", exc)
            return "Failed to rename the folder."

    def copy_folder(self, source_folder_name: str, destination_folder_name: str) -> str:
        if not source_folder_name or not destination_folder_name:
            return "Source or destination folder name not provided."

        try:
            source_path = self._build_path(source_folder_name)
            destination_path = self._build_path(destination_folder_name)
            shutil.copytree(source_path, destination_path, dirs_exist_ok=True)
            return f"Folder '{source_folder_name}' copied to '{destination_folder_name}' successfully."
        except FileNotFoundError:
            return f"Folder '{source_folder_name}' not found."
        except Exception as exc:
            logging.error("Folder Copy Error: %s", exc)
            return "Failed to copy the folder."

    def zip_folder(self, folder_name: str) -> str:
        if not folder_name:
            return "No folder name provided."

        try:
            folder_path = self._build_path(folder_name)
            archive_path = self._build_path(f"{folder_name}.zip")
            with zipfile.ZipFile(archive_path, "w", zipfile.ZIP_DEFLATED) as archive:
                for file_path in folder_path.rglob("*"):
                    if file_path.is_file():
                        archive.write(file_path, file_path.relative_to(folder_path))
            return f"Folder '{folder_name}' zipped successfully."
        except FileNotFoundError:
            return f"Folder '{folder_name}' not found."
        except Exception as exc:
            logging.error("Zip Error: %s", exc)
            return "Failed to zip the folder."
